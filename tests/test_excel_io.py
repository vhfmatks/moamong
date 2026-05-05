from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from moamong_app.excel_io import (
    REQUIRED_PRODUCT_COLUMNS,
    WorkbookData,
    export_results,
    load_category_workbook,
    load_product_workbook,
)
from moamong_app.models import Category, GeneratedRow, RowProcessResult, RowStatus


def make_product_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "List1(1-2)"
    headers = list(REQUIRED_PRODUCT_COLUMNS) + ["등록일"]
    ws.append(headers)
    ws.append(["P001", "얼굴배게", "마사지샵 얼굴 쿠션 베개", "", "마사지베개", "WB100", "2026-05-05"])
    ws.append(["P002", "가위", "수초 가위", "", "수초가위", "WB200", "2026-05-05"])
    wb.save(path)


def make_category_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mycate1(1-2)"
    ws.append(["마이카테", "마이카테명", "등록일"])
    ws.append(["WB100", "생활용품 > 침구 > 베개 >", "2026-05-05"])
    ws.append(["WB200", "반려동물 > 관상어용품 > 수초관리 >", "2026-05-05"])
    wb.save(path)


def test_load_product_workbook_reads_rows(tmp_path: Path) -> None:
    path = tmp_path / "product.xlsx"
    make_product_workbook(path)

    data = load_product_workbook(path)

    assert isinstance(data, WorkbookData)
    assert data.sheet_name == "List1(1-2)"
    assert len(data.rows) == 2
    assert data.rows[0].text("원본상품명(참고용)") == "마사지샵 얼굴 쿠션 베개"


def test_load_product_workbook_reports_row_progress(tmp_path: Path) -> None:
    path = tmp_path / "product.xlsx"
    make_product_workbook(path)
    events: list[tuple[int, int]] = []

    load_product_workbook(path, progress_callback=lambda current, total: events.append((current, total)))

    assert events[0] == (0, 2)
    assert events[-1] == (2, 2)


def test_load_category_workbook_reads_mycate_columns(tmp_path: Path) -> None:
    path = tmp_path / "category.xlsx"
    make_category_workbook(path)

    categories = load_category_workbook(path)

    assert categories == [
        Category(mycate="WB100", name="생활용품 > 침구 > 베개 >"),
        Category(mycate="WB200", name="반려동물 > 관상어용품 > 수초관리 >"),
    ]


def test_load_category_workbook_reports_row_progress(tmp_path: Path) -> None:
    path = tmp_path / "category.xlsx"
    make_category_workbook(path)
    events: list[tuple[int, int]] = []

    load_category_workbook(path, progress_callback=lambda current, total: events.append((current, total)))

    assert events[0] == (0, 2)
    assert events[-1] == (2, 2)


def test_export_results_updates_columns_and_adds_audit_columns(tmp_path: Path) -> None:
    product_path = tmp_path / "product.xlsx"
    output_path = tmp_path / "product_mapped.xlsx"
    make_product_workbook(product_path)
    data = load_product_workbook(product_path)
    result = RowProcessResult(
        row_index=0,
        status=RowStatus.DONE,
        generated=GeneratedRow(
            product_name="얼굴전용 마사지 베개",
            keywords=["마사지베개", "얼굴쿠션"],
            mycate="WB100",
            mycate_name="생활용품 > 침구 > 베개 >",
            confidence=0.91,
            review_reason="",
        ),
    )

    export_results(data, [result], output_path)

    wb = load_workbook(output_path)
    ws = wb["List1(1-2)"]
    headers = [cell.value for cell in ws[1]]
    row = [cell.value for cell in ws[2]]
    values = dict(zip(headers, row))
    assert values["상품명"] == "얼굴전용 마사지 베개"
    assert values["키워드"] == "마사지베개,얼굴쿠션"
    assert values["마이카테"] == "WB100"
    assert values["LLM_status"] == "완료"
    assert values["LLM_마이카테명"] == "생활용품 > 침구 > 베개 >"


def test_export_product_keyword_results_preserves_mycate_and_writes_site_keywords(tmp_path: Path) -> None:
    product_path = tmp_path / "product.xlsx"
    output_path = tmp_path / "product_keywords.xlsx"
    make_product_workbook(product_path)
    data = load_product_workbook(product_path)
    result = RowProcessResult(
        row_index=0,
        status=RowStatus.DONE,
        generated=GeneratedRow(
            product_name="얼굴전용 마사지 베개",
            keywords=["얼굴베개", "마사지얼굴베개"],
        ),
        search_keyword="얼굴 쿠션",
        site_keywords={
            "coupang": ["얼굴베개", "마사지얼굴쿠션"],
            "naver": ["안면베개"],
        },
        step_statuses={
            "search_keyword": "완료",
            "coupang": "완료",
            "naver": "완료",
            "llm": "완료",
        },
    )

    export_results(data, [result], output_path)

    wb = load_workbook(output_path)
    ws = wb["List1(1-2)"]
    headers = [cell.value for cell in ws[1]]
    row = [cell.value for cell in ws[2]]
    values = dict(zip(headers, row))
    assert values["상품명"] == "얼굴전용 마사지 베개"
    assert values["키워드"] == "얼굴베개,마사지얼굴베개"
    assert values["마이카테"] == "WB100"
    assert values["LLM_검색어"] == "얼굴 쿠션"
    assert values["쿠팡_키워드"] == "얼굴베개,마사지얼굴쿠션"
    assert values["네이버_키워드"] == "안면베개"
    assert values["쿠팡_status"] == "완료"
    assert values["네이버_status"] == "완료"
    assert values["검색어_status"] == "완료"
    assert values["LLM_step_status"] == "완료"
