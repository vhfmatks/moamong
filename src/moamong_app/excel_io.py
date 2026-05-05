from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import load_workbook

from moamong_app.keyword_fetcher import (
    SITE_KEYWORD_COLUMNS,
    SITE_STATUS_COLUMNS,
    site_keyword_column,
    site_status_column,
)
from moamong_app.models import Category, ProductRow, RowProcessResult


REQUIRED_PRODUCT_COLUMNS = (
    "상품코드",
    "상품명",
    "원본상품명(참고용)",
    "옵션명",
    "키워드",
    "마이카테",
)

ProgressCallback = Callable[[int, int], None]


@dataclass
class WorkbookData:
    path: Path
    sheet_name: str
    rows: list[ProductRow]
    columns: list[str]


def _first_non_empty_sheet(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            first_row = next(ws.iter_rows(max_row=1, values_only=True), None)
            if first_row is not None and any(_cell_text(value) for value in first_row):
                return sheet_name
        return wb.sheetnames[0]
    finally:
        wb.close()


def _product_sheet_name(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if sheet_name.startswith("List1"):
                return sheet_name
    finally:
        wb.close()
    return _first_non_empty_sheet(path)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _progress(progress_callback: ProgressCallback | None, current: int, total: int) -> None:
    if progress_callback is not None:
        progress_callback(current, total)


def load_product_workbook(
    path: str | Path,
    progress_callback: ProgressCallback | None = None,
) -> WorkbookData:
    workbook_path = Path(path)
    sheet_name = _product_sheet_name(workbook_path)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("Product workbook has no header row")
        columns = [_cell_text(column) for column in header_row]
        missing = [column for column in REQUIRED_PRODUCT_COLUMNS if column not in columns]
        if missing:
            raise ValueError(f"Product workbook missing required columns: {', '.join(missing)}")

        total = max((ws.max_row or 1) - 1, 0)
        _progress(progress_callback, 0, total)
        rows: list[ProductRow] = []
        for index, row_values in enumerate(iterator):
            values = {
                column: _cell_text(row_values[column_index]) if column_index < len(row_values) else ""
                for column_index, column in enumerate(columns)
            }
            rows.append(ProductRow(index=index, values=values))
            _progress(progress_callback, index + 1, total)
        return WorkbookData(path=workbook_path, sheet_name=sheet_name, rows=rows, columns=columns)
    finally:
        wb.close()


def load_category_workbook(
    path: str | Path,
    progress_callback: ProgressCallback | None = None,
) -> list[Category]:
    workbook_path = Path(path)
    sheet_name = _first_non_empty_sheet(workbook_path)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("Category workbook has no header row")
        columns = [_cell_text(column) for column in header_row]
        missing = [column for column in ("마이카테", "마이카테명") if column not in columns]
        if missing:
            raise ValueError(f"Category workbook missing required columns: {', '.join(missing)}")
        mycate_index = columns.index("마이카테")
        name_index = columns.index("마이카테명")
        total = max((ws.max_row or 1) - 1, 0)
        _progress(progress_callback, 0, total)

        categories: list[Category] = []
        for current, row_values in enumerate(iterator, start=1):
            mycate = _cell_text(row_values[mycate_index]) if mycate_index < len(row_values) else ""
            name = _cell_text(row_values[name_index]) if name_index < len(row_values) else ""
            if mycate and name:
                categories.append(Category(mycate=mycate, name=name))
            _progress(progress_callback, current, total)
        if not categories:
            raise ValueError("Category workbook has no usable category rows")
        return categories
    finally:
        wb.close()


def _ensure_headers(ws, headers: Iterable[str]) -> dict[str, int]:
    existing = [cell.value for cell in ws[1]]
    for header in headers:
        if header not in existing:
            ws.cell(row=1, column=len(existing) + 1, value=header)
            existing.append(header)
    return {str(header): idx + 1 for idx, header in enumerate(existing)}


def _site_keyword_headers(results: list[RowProcessResult]) -> list[str]:
    headers = list(SITE_KEYWORD_COLUMNS)
    for result in results:
        for site in result.site_keywords:
            header = site_keyword_column(site)
            if header not in headers:
                headers.append(header)
    return headers


def _step_status_headers(results: list[RowProcessResult]) -> list[str]:
    headers = ["검색어_status", *SITE_STATUS_COLUMNS, "LLM_step_status"]
    for result in results:
        for step in result.step_statuses:
            if step == "llm":
                header = "LLM_step_status"
            elif step == "search_keyword":
                header = "검색어_status"
            else:
                header = site_status_column(step)
            if header not in headers:
                headers.append(header)
    return headers


def export_results(data: WorkbookData, results: list[RowProcessResult], output_path: str | Path) -> Path:
    output = Path(output_path)
    if output.resolve() == data.path.resolve():
        raise ValueError("Output path must be different from the source workbook path")
    wb = load_workbook(data.path)
    ws = wb[data.sheet_name]
    headers = _ensure_headers(
        ws,
        [
            *_site_keyword_headers(results),
            *_step_status_headers(results),
            "LLM_검색어",
            "LLM_상품명",
            "LLM_키워드",
            "LLM_마이카테",
            "LLM_마이카테명",
            "LLM_confidence",
            "LLM_status",
            "LLM_review_reason",
        ],
    )
    for result in results:
        if result.row_index < 0 or result.row_index >= len(data.rows):
            raise ValueError(f"Result row index out of range: {result.row_index}")
        excel_row = result.row_index + 2
        generated = result.generated
        ws.cell(excel_row, headers["LLM_status"], result.status.value)
        ws.cell(excel_row, headers["LLM_review_reason"], result.message)
        ws.cell(excel_row, headers["LLM_검색어"], result.search_keyword)
        for site, keywords in result.site_keywords.items():
            header = site_keyword_column(site)
            value = ",".join(keyword.strip() for keyword in keywords if keyword.strip())
            if not value and site in result.site_errors:
                value = f"오류: {result.site_errors[site]}"
            ws.cell(excel_row, headers[header], value)
        for step, status in result.step_statuses.items():
            if step == "llm":
                header = "LLM_step_status"
            elif step == "search_keyword":
                header = "검색어_status"
            else:
                header = site_status_column(step)
            ws.cell(excel_row, headers[header], status)
        if generated is None:
            continue
        ws.cell(excel_row, headers["상품명"], generated.product_name)
        ws.cell(excel_row, headers["키워드"], generated.keywords_text)
        if generated.mycate.strip():
            ws.cell(excel_row, headers["마이카테"], generated.mycate)
        ws.cell(excel_row, headers["LLM_상품명"], generated.product_name)
        ws.cell(excel_row, headers["LLM_키워드"], generated.keywords_text)
        ws.cell(excel_row, headers["LLM_마이카테"], generated.mycate)
        ws.cell(excel_row, headers["LLM_마이카테명"], generated.mycate_name)
        ws.cell(excel_row, headers["LLM_confidence"], generated.confidence)
        ws.cell(excel_row, headers["LLM_review_reason"], generated.review_reason or result.message)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
